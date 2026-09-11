import io
import logging

from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from PyPDF2 import PdfReader
from sqlalchemy.ext.asyncio import AsyncSession
from tqdm import tqdm

from db.db_models import Document, DocumentChunk, Session
from db.vectordb import vectordb
from services.llm_service import llm_service

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = (".txt", ".md", ".docx", ".pdf")


class IngestionService:

    async def ingest(self, file, db: AsyncSession, session_id: str):
        logger.info(f"Starting ingestion for file: {file.filename}")

        # 1. Extract raw text
        raw_text = await self._extract_text(file)
        logger.info(f"Extracted {len(raw_text)} characters from uploaded file")

        if not raw_text.strip():
            raise ValueError(
                f"Could not extract any text from '{file.filename}'. "
                f"Supported formats: {', '.join(SUPPORTED_EXTENSIONS)}"
            )

        # 2. Clean text
        cleaned = self._clean_text(raw_text)
        logger.info(f"Cleaned text length: {len(cleaned)}")

        # 3. Chunk text
        chunks = self._chunk_text(cleaned)
        logger.info(f"Generated {len(chunks)} chunks")

        # 4. Add Document row to SQLite
        doc = Document(filename=file.filename, content_type=file.content_type, session_id=session_id)
        db.add(doc)
        await db.flush()  # populates doc.id before we need it below
        doc_id = doc.id
        logger.info(f"Added document to Document table: {doc_id}")

        # 5. Save chunks + embeddings into SQLite + ChromaDB
        logger.info("Creating chunk embeddings and adding them to ChromaDB")
        for idx, chunk_text in enumerate(tqdm(chunks, desc="Embedding chunks", unit="chunk")):
            chunk = DocumentChunk(document_id=doc_id, session_id=session_id, chunk_index=idx, text=chunk_text)
            db.add(chunk)
            await db.flush()
            chunk_id = chunk.id

            emb = await llm_service.embed(chunk_text)

            metadata = {
                "session_id": session_id,
                "doc_id": doc_id,
                "chunk_id": chunk_id,
                "chunk_index": idx,
                "text": chunk_text,
            }
            vector_id = f"{session_id}_{doc_id}_{chunk_id}"
            vectordb.add_vector(collection_name="chunks", embedding=emb, metadata=metadata, vector_id=vector_id)

        logger.info(f"Completed ChromaDB ingestion: {len(chunks)} vectors added")

        # 6. Generate a friendly session title from the first chunk
        try:
            preview_text = chunks[0][:500]
            session_name = await llm_service.generate_session_title(preview_text)
            session_row = await db.get(Session, session_id)
            session_row.session_name = session_name
            logger.info(f"[{session_id}] Generated session name: {session_name}")
        except Exception as e:
            logger.error(f"Failed to generate session name: {e}")

        await db.commit()
        logger.info(f"Finished ingestion for document: {doc_id} ({len(chunks)} chunks)")

        return doc_id

    async def _extract_text(self, file) -> str:
        """Extract text from .pdf, .txt, .docx, .md files."""
        content = await file.read()
        filename = file.filename.lower()

        if filename.endswith(".txt") or filename.endswith(".md"):
            logger.info("Detected TXT/MD file")
            return content.decode("utf-8", errors="ignore")

        if filename.endswith(".docx"):
            logger.info("Detected DOCX file")
            doc = DocxDocument(io.BytesIO(content))
            return "\n".join(par.text for par in doc.paragraphs)

        if filename.endswith(".pdf"):
            logger.info("Detected PDF file")
            reader = PdfReader(io.BytesIO(content))
            extracted = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(extracted)

        logger.warning(f"Unsupported file type for '{file.filename}'")
        return ""

    def _clean_text(self, text: str) -> str:
        return text.replace("\r", "").strip()

    def _chunk_text(self, text: str) -> list[str]:
        splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=150, length_function=len)
        return splitter.split_text(text)


ingestion_service = IngestionService()
